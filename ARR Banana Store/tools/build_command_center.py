"""Builds ARR_Command_Center.xlsx: one workbook for daily input and the profit dashboard.

Usage: python3 build_command_center.py <monthly_records.xlsx> <out.xlsx>
The first argument is an .xlsx export of "Monthly Records Report - 24/25/26"; it seeds
the History tab and the October 2026 test rows.
"""
import sys
import datetime as dt

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

SRC, OUT = sys.argv[1], sys.argv[2]

PESO = '"₱"#,##0;("₱"#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = 'd mmm yyyy'
MONTH = 'mmm yyyy'
BLUE = Font(name='Arial', color='0000FF')
BOLD = Font(name='Arial', bold=True)
TITLE = Font(name='Arial', bold=True, size=14)
HEAD_FILL = PatternFill('solid', fgColor='F4D03F')
INPUT_FILL = PatternFill('solid', fgColor='FFFF00')
BRANCHES = ['Main', 'Rizal']
CATEGORIES = ['Rent', 'Salary (net paid)', 'Cash Advance', 'Food & Water', 'Utilities',
              'Plastic', 'Trash', 'Repairs & Maintenance', 'Ads & Tarpaulin',
              'Trucking & Labor', 'Other']

wb = openpyxl.Workbook()


def header(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical='center')
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def title(ws, text, legend):
    ws['A1'] = text
    ws['A1'].font = TITLE
    ws['A2'] = legend
    ws['A2'].font = Font(name='Arial', italic=True, color='666666')


def put(ws, row, values, fmts):
    for i, (v, f) in enumerate(zip(values, fmts), start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = BLUE if not (isinstance(v, str) and v.startswith('=')) else Font(name='Arial')
        if f:
            c.number_format = f


def dropdown(ws, formula, rng):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


# ---------------------------------------------------------------- Settings
st = wb.active
st.title = 'Settings'
title(st, 'Settings', 'Blue cells are inputs. Change a threshold here and every alert uses it.')
st['A4'], st['B4'] = 'Branches', None
st['A4'].font = BOLD
for i, b in enumerate(BRANCHES):
    st.cell(row=5 + i, column=1, value=b).font = BLUE
st['C4'] = 'Expense categories'
st['C4'].font = BOLD
for i, c in enumerate(CATEGORIES):
    st.cell(row=5 + i, column=3, value=c).font = BLUE
st['E4'] = 'Alert thresholds'
st['E4'].font = BOLD
st['E5'], st['F5'] = 'Cash short alert, per day (₱)', 300
st['E6'], st['F6'] = 'Spoilage alert, % of sales', 0.05
for r, f in ((5, PESO), (6, PCT)):
    st.cell(row=r, column=6).font = BLUE
    st.cell(row=r, column=6).number_format = f
st['F5'].comment = Comment('Source: user-provided assumption, 6 Oct 2026. Change as needed.', 'ARR')
st['F6'].comment = Comment('Source: user-provided assumption, 6 Oct 2026. Recent spoilage ran about 4.6% of purchases.', 'ARR')
for col, w in (('A', 14), ('B', 4), ('C', 24), ('D', 4), ('E', 30), ('F', 12)):
    st.column_dimensions[col].width = w
BR_LIST = 'Settings!$A$5:$A$6'
CAT_LIST = f'Settings!$C$5:$C${4 + len(CATEGORIES)}'

# ---------------------------------------------------------------- Daily Sales
ds = wb.create_sheet('Daily Sales')
title(ds, 'Daily sales', 'One row per branch per day. Blue = typed in. Cash short is negative, surplus positive.')
header(ds, 4, ['Date', 'Branch', 'Retail sales (₱)', 'Wholesale sales (₱)', 'Cash short (-) / over (+) (₱)',
               'Spoilage at cost (₱)', 'Lakatan sold (kg)', 'Lakatan spoiled (kg)', 'Notes'],
       [13, 10, 14, 14, 16, 14, 12, 12, 30])
dropdown(ds, BR_LIST, 'B5:B2000')

src = openpyxl.load_workbook(SRC, data_only=True)
od = src['October - Dailies26']
row = 5
for r in od.iter_rows(min_row=6, max_row=40, values_only=True):
    d, main, rizal, whole, var_m, var_r = r[0], r[3], r[4], r[5], r[6], r[7]
    if not isinstance(d, dt.datetime) or main in (None, ''):
        continue
    for br, sales, wh, var in (('Main', main, whole or 0, var_m), ('Rizal', rizal, 0, var_r)):
        put(ds, row, [d, br, sales, wh, var, None, None, None, None],
            [DATE, None, PESO, PESO, PESO, PESO, '0.0', '0.0', None])
        row += 1
ds['A5'].comment = Comment('Source: Monthly Records Report - 24/25/26, tab "October - Dailies26", '
                           'rows 1-6 Oct 2026, read 6 Oct 2026. Wholesale was assigned to Main. '
                           'Spoilage and kg columns were not in that tab.', 'ARR')

# ---------------------------------------------------------------- Deliveries
dl = wb.create_sheet('Deliveries')
title(dl, 'Fruit deliveries', 'One row per item per delivery. Trucking goes in Expenses as "Trucking & Labor".')
header(dl, 4, ['Date', 'Supplier', 'Item', 'Kg', 'Amount paid (₱)', 'Notes'],
       [13, 14, 14, 10, 15, 40])
oct_purch = next(r[4] for r in src['DASHBOARD'].iter_rows(min_row=4, values_only=True)
                 if isinstance(r[0], dt.datetime) and r[0] == dt.datetime(2026, 10, 1))
put(dl, 5, [dt.datetime(2026, 10, 6), 'Various', 'All fruit', None, oct_purch,
            '1-6 Oct total, itemized kg not available'],
    [DATE, None, None, '#,##0.0', PESO, None])
dl['E5'].comment = Comment('Source: Monthly Records Report - 24/25/26, tab "DASHBOARD", '
                           'Inventory Purchase for Oct 2026, read 6 Oct 2026.', 'ARR')

# ---------------------------------------------------------------- Expenses
ex = wb.create_sheet('Expenses')
title(ex, 'Expenses', 'Record cash advances when given (category "Cash Advance") and salary as the net amount paid.')
header(ex, 4, ['Date', 'Branch', 'Category', 'Description', 'Amount (₱)', 'Paid from'],
       [13, 10, 22, 40, 13, 14])
dropdown(ex, BR_LIST, 'B5:B3000')
dropdown(ex, CAT_LIST, 'C5:C3000')
dropdown(ex, '"Drawer,Owner cash,GCash"', 'F5:F3000')
oct_exp = [('Main', 'Plastic', 57), ('Main', 'Trash', 20), ('Main', 'Food & Water', 2401),
           ('Main', 'Trucking & Labor', 16612), ('Rizal', 'Food & Water', 395),
           ('Main', 'Other', 4025), ('Rizal', 'Ads & Tarpaulin', 10)]
for i, (br, cat, amt) in enumerate(oct_exp):
    put(ex, 5 + i, [dt.datetime(2026, 10, 6), br, cat, '1-6 Oct total from old expense sheet', amt, None],
        [DATE, None, None, None, PESO, None])
ex['A5'].comment = Comment('Source: [Main] and [Branch1] ARR BANANA Monthly-Expenses - 2026, tab "OCT", '
                           'category totals for 1-6 Oct 2026, read 6 Oct 2026.', 'ARR')

# ---------------------------------------------------------------- Stock Count
sc = wb.create_sheet('Stock Count')
title(sc, 'Month-end stock count', 'Count what is left on the last day of each month. "Month" = the month that just ended (first day).')
header(sc, 4, ['Month', 'Branch', 'Item', 'Kg on hand', 'Cost per kg (₱)'],
       [13, 10, 14, 12, 14])
dropdown(sc, BR_LIST, 'B5:B1000')

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet('Dashboard', 0)
db['A1'] = 'ARR Banana Store dashboard'
db['A1'].font = TITLE
db['A2'] = 'Change the month in the yellow cell. Everything else is calculated from the input tabs.'
db['A2'].font = Font(name='Arial', italic=True, color='666666')
db['A3'] = 'Month'
db['A3'].font = BOLD
db['B3'] = dt.datetime(2026, 10, 1)
db['B3'].number_format = MONTH
db['B3'].font = BLUE
db['B3'].fill = INPUT_FILL
for col, w in (('A', 46), ('B', 15), ('C', 15), ('D', 15)):
    db.column_dimensions[col].width = w

M = '">="&$B$3'
N = '"<"&EDATE($B$3,1)'


def dated(sheet):
    return f"'{sheet}'!$A:$A,{M},'{sheet}'!$A:$A,{N}"


def per_branch(sheet, col, extra=''):
    return [f"=SUMIFS('{sheet}'!${col}:${col},'{sheet}'!$B:$B,{L}$5,{dated(sheet)}{extra})"
            for L in ('B', 'C')]


db['A5'], db['B5'], db['C5'], db['D5'] = 'Sales', 'Main', 'Rizal', 'Total'
layout = [
    (6, 'Retail sales', per_branch('Daily Sales', 'C'), 'sum', PESO),
    (7, 'Wholesale sales', per_branch('Daily Sales', 'D'), 'sum', PESO),
    (8, 'Total sales', ['=B6+B7', '=C6+C7'], 'sum', PESO),
    (9, 'Spoilage at cost', per_branch('Daily Sales', 'F'), 'sum', PESO),
    (10, 'Spoilage, % of sales', ['=IF(B8=0,0,B9/B8)', '=IF(C8=0,0,C9/C8)'], '=IF(D8=0,0,D9/D8)', PCT),
    (11, 'Cash short (-) / over (+)', per_branch('Daily Sales', 'E'), 'sum', PESO),
    (12, 'Days reported', [f"=COUNTIFS('Daily Sales'!$B:$B,{L}$5,{dated('Daily Sales')})" for L in 'BC'], None, '0'),
]
db['A14'], db['D14'] = 'Cost of fruit (whole business)', 'Total'
layout += [
    (15, 'Fruit purchases', None, f"=SUMIFS(Deliveries!$E:$E,{dated('Deliveries')})", PESO),
    (16, 'Trucking & labor', None, f"=SUMIFS(Expenses!$E:$E,Expenses!$C:$C,\"Trucking & Labor\",{dated('Expenses')})", PESO),
    (17, 'Stock at start of month', None, "=SUMPRODUCT(('Stock Count'!$A$5:$A$5000=EDATE($B$3,-1))*'Stock Count'!$D$5:$D$5000*'Stock Count'!$E$5:$E$5000)", PESO),
    (18, 'Stock at end of month', None, "=SUMPRODUCT(('Stock Count'!$A$5:$A$5000=$B$3)*'Stock Count'!$D$5:$D$5000*'Stock Count'!$E$5:$E$5000)", PESO),
    (19, 'Cost of fruit sold', None, '=D15+D16+D17-D18', PESO),
    (20, 'Gross margin, %', None, '=IF(D8=0,0,(D8-D19)/D8)', PCT),
    (21, 'Kg delivered (rows with kg filled in)', None, f"=SUMIFS(Deliveries!$D:$D,{dated('Deliveries')})", '#,##0'),
]
db['A22'], db['B22'], db['C22'], db['D22'] = 'Profit', 'Main', 'Rizal', 'Total'
opex = lambda L: (f"=SUMIFS(Expenses!$E:$E,Expenses!$B:$B,{L}$22,{dated('Expenses')})"
                  f"-SUMIFS(Expenses!$E:$E,Expenses!$B:$B,{L}$22,Expenses!$C:$C,\"Trucking & Labor\",{dated('Expenses')})")
staff = lambda L: (f"=SUMIFS(Expenses!$E:$E,Expenses!$B:$B,{L}$22,Expenses!$C:$C,\"Salary (net paid)\",{dated('Expenses')})"
                   f"+SUMIFS(Expenses!$E:$E,Expenses!$B:$B,{L}$22,Expenses!$C:$C,\"Cash Advance\",{dated('Expenses')})")
layout += [
    (23, 'Gross profit (branch split by share of sales)', ['=B8*$D$20', '=C8*$D$20'], '=D8-D19', PESO),
    (24, 'Operating expenses', [opex('B'), opex('C')], 'sum', PESO),
    (25, '   of which staff (salary + cash advances)', [staff('B'), staff('C')], 'sum', PESO),
    (26, 'Net profit', ['=B23-B24', '=C23-C24'], '=D23-D24', PESO),
    (27, 'Net margin, %', ['=IF(B8=0,0,B26/B8)', '=IF(C8=0,0,C26/C8)'], '=IF(D8=0,0,D26/D8)', PCT),
]
for r, label, br, tot, fmt in layout:
    db.cell(row=r, column=1, value=label)
    if br:
        for j, f in enumerate(br):
            db.cell(row=r, column=2 + j, value=f).number_format = fmt
    if tot == 'sum':
        tot = f'=SUM(B{r}:C{r})'
    if tot:
        db.cell(row=r, column=4, value=tot).number_format = fmt
for r in (5, 14, 22):
    for c in range(1, 5):
        db.cell(row=r, column=c).font = BOLD
        db.cell(row=r, column=c).fill = HEAD_FILL
for r in (8, 19, 26):
    for c in range(1, 5):
        db.cell(row=r, column=c).font = BOLD

db['A29'], db['B29'], db['C29'] = 'Alerts', 'Main', 'Rizal'
for c in range(1, 4):
    db.cell(row=29, column=c).font = BOLD
    db.cell(row=29, column=c).fill = HEAD_FILL
db['A30'] = 'Days missing a report (so far this month)'
db['A31'] = 'Days cash was short by more than the limit'
db['A32'] = 'Spoilage above the limit'
db['A33'] = 'Stock count for this month entered'
db['A35'] = 'Days so far this month'
db['B35'] = '=MAX(0,MIN(TODAY(),EOMONTH($B$3,0))-$B$3+1)'
for L in 'BC':
    db[f'{L}30'] = f'=MAX(0,$B$35-{L}12)'
    db[f'{L}31'] = (f"=COUNTIFS('Daily Sales'!$B:$B,{L}$29,{dated('Daily Sales')},"
                    f"'Daily Sales'!$E:$E,\"<\"&-Settings!$F$5)")
    db[f'{L}32'] = f'=IF({L}10>Settings!$F$6,"CHECK","OK")'
db['B33'] = ("=IF(COUNTIFS('Stock Count'!$A:$A,$B$3)>0,\"Yes\","
             "\"No: until counted, profit treats all fruit bought this month as sold\")")
db['A37'] = ('How profit is worked out: cost of fruit sold = purchases + trucking + stock at start - stock at end. '
             'Branch gross profit is estimated by each branch\'s share of sales, because deliveries are shared.')
db['A37'].alignment = Alignment(wrap_text=True)
db.merge_cells('A37:D38')
db.row_dimensions[37].height = 30

# Conditional colors on alerts
from openpyxl.formatting.rule import FormulaRule
red = PatternFill('solid', fgColor='F5B7B1')
db.conditional_formatting.add('B30:C31', FormulaRule(formula=['B30>0'], fill=red))
db.conditional_formatting.add('B32:C32', FormulaRule(formula=['B32="CHECK"'], fill=red))
db.conditional_formatting.add('B33', FormulaRule(formula=['LEFT(B33,2)="No"'], fill=red))
db.conditional_formatting.add('B26:D26', FormulaRule(formula=['B26<0'], fill=red))

# ---------------------------------------------------------------- Monthly Trend
tr = wb.create_sheet('Monthly Trend', 1)
title(tr, 'Monthly trend, 2026', 'Calculated from the input tabs. Months before October 2026 are on the History tab.')
tr['A4'] = 'Month'
tr['A4'].font = BOLD
tr.column_dimensions['A'].width = 26
for i in range(12):
    L = openpyxl.utils.get_column_letter(2 + i)
    c = tr.cell(row=4, column=2 + i, value=dt.datetime(2026, 1 + i, 1))
    c.number_format = 'mmm'
    c.font = BOLD
    c.fill = HEAD_FILL
    tr.column_dimensions[L].width = 11
    m, n = f'">="&{L}$4', f'"<"&EDATE({L}$4,1)'
    d = lambda s: f"'{s}'!$A:$A,{m},'{s}'!$A:$A,{n}"
    f = {
        5: f"=SUMIFS('Daily Sales'!$C:$C,{d('Daily Sales')})+SUMIFS('Daily Sales'!$D:$D,{d('Daily Sales')})",
        6: f"=SUMIFS(Deliveries!$E:$E,{d('Deliveries')})+SUMIFS(Expenses!$E:$E,Expenses!$C:$C,\"Trucking & Labor\",{d('Expenses')})"
           f"+SUMPRODUCT(('Stock Count'!$A$5:$A$5000=EDATE({L}$4,-1))*'Stock Count'!$D$5:$D$5000*'Stock Count'!$E$5:$E$5000)"
           f"-SUMPRODUCT(('Stock Count'!$A$5:$A$5000={L}$4)*'Stock Count'!$D$5:$D$5000*'Stock Count'!$E$5:$E$5000)",
        7: f'={L}5-{L}6',
        8: f"=SUMIFS(Expenses!$E:$E,{d('Expenses')})-SUMIFS(Expenses!$E:$E,Expenses!$C:$C,\"Trucking & Labor\",{d('Expenses')})",
        9: f'={L}7-{L}8',
        10: f"=SUMIFS('Daily Sales'!$F:$F,{d('Daily Sales')})",
    }
    for r, formula in f.items():
        tr.cell(row=r, column=2 + i, value=formula).number_format = PESO
for r, label in ((5, 'Total sales'), (6, 'Cost of fruit sold'), (7, 'Gross profit'),
                 (8, 'Operating expenses'), (9, 'Net profit'), (10, 'Spoilage at cost')):
    tr.cell(row=r, column=1, value=label)
for c in range(1, 14):
    tr.cell(row=9, column=c).font = BOLD
tr.freeze_panes = 'B5'

# ---------------------------------------------------------------- History
hi = wb.create_sheet('History')
title(hi, 'History, Sep 2024 to Sep 2026', 'Copied from the old Monthly Records dashboard. "Cash result" is sales minus cash spent that month, not profit.')
header(hi, 4, ['Month', 'Sales (₱)', 'Fruit purchases (₱)', 'Expenses (₱)', 'Cash result (₱)'],
       [12, 15, 17, 15, 15])
row = 5
for r in src['DASHBOARD'].iter_rows(min_row=4, values_only=True):
    if isinstance(r[0], dt.datetime) and r[0] < dt.datetime(2026, 10, 1) and isinstance(r[3], (int, float)):
        put(hi, row, [r[0], r[3], r[4], r[5], f'=B{row}-C{row}-D{row}'], [MONTH, PESO, PESO, PESO, PESO])
        row += 1
hi['A4'].comment = Comment('Source: Monthly Records Report - 24/25/26, tab "DASHBOARD", columns Sales, '
                           'Inventory Purchase, Expenses, Sep 2024 to Sep 2026, read 6 Oct 2026.', 'ARR')
put(hi, row, ['Total', f'=SUM(B5:B{row-1})', f'=SUM(C5:C{row-1})', f'=SUM(D5:D{row-1})', f'=SUM(E5:E{row-1})'],
    [None, PESO, PESO, PESO, PESO])
for c in range(1, 6):
    hi.cell(row=row, column=c).font = BOLD

for ws in wb.worksheets:
    for r in ws.iter_rows():
        for c in r:
            if c.font and c.font.name != 'Arial':
                c.font = Font(name='Arial', bold=c.font.bold, italic=c.font.italic,
                              color=c.font.color, size=c.font.size)

wb.save(OUT)
print('saved', OUT)
